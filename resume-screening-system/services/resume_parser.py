import re
import os
import json

COMMON_SKILLS = [
    # Languages
    "Python", "JavaScript", "TypeScript", "Java", "C", "C++", "C#", "Ruby", "Go", "Golang",
    "Rust", "PHP", "Swift", "Kotlin", "Dart", "R", "Scala", "Shell", "Bash", "PowerShell",
    "HTML", "HTML5", "CSS", "CSS3", "Sass", "SCSS",
    
    # Frameworks & Libraries
    "React", "React.js", "Angular", "Vue", "Vue.js", "Next.js", "Node.js", "Express", "Express.js",
    "Django", "Flask", "FastAPI", "Spring", "Spring Boot", "ASP.NET", ".NET Core", "Laravel",
    "Bootstrap", "TailwindCSS", "Tailwind", "jQuery", "Redux", "GraphQL", "REST APIs", "RESTful APIs",
    
    # Databases & Caching
    "SQL", "MySQL", "PostgreSQL", "SQLite", "MongoDB", "Redis", "Oracle", "Cassandra",
    "DynamoDB", "Elasticsearch", "Firebase", "Supabase", "MariaDB",
    
    # Cloud & DevOps
    "AWS", "Amazon Web Services", "Azure", "Microsoft Azure", "GCP", "Google Cloud",
    "Docker", "Kubernetes", "Git", "GitHub", "GitLab", "CI/CD", "Jenkins", "Terraform",
    "Ansible", "Linux", "Unix", "Nginx", "Apache",
    
    # Data Science & AI/ML
    "Machine Learning", "Deep Learning", "Artificial Intelligence", "NLP", "Natural Language Processing",
    "Computer Vision", "scikit-learn", "TensorFlow", "Keras", "PyTorch", "Pandas", "NumPy",
    "Matplotlib", "Seaborn", "Data Analysis", "Data Visualization", "OpenCV", "Tableau",
    "Power BI", "Spark", "Hadoop", "Kafka",
    
    # Testing & Methodologies
    "Agile", "Scrum", "Jira", "Unit Testing", "PyTest", "Selenium", "Jest", "Mocha",
    "TDD", "Microservices", "Design Patterns", "Software Architecture", "UI/UX", "Figma"
]

def extract_text_from_pdf(filepath):
    text = ""
    try:
        from pypdf import PdfReader
        reader = PdfReader(filepath)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"
    except Exception as e:
        print(f"pypdf error: {e}")
        try:
            import pdfplumber
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text += t + "\n"
        except Exception as e2:
            print(f"pdfplumber error: {e2}")
    return text.strip()

def extract_text_from_docx(filepath):
    text = ""
    try:
        import docx
        doc = docx.Document(filepath)
        for p in doc.paragraphs:
            if p.text:
                text += p.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text += " | ".join(row_text) + "\n"
    except Exception as e:
        print(f"docx error: {e}")
    return text.strip()

def extract_text_from_txt(filepath):
    try:
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read().strip()
    except Exception:
        return ""

def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.pdf':
        return extract_text_from_pdf(filepath)
    elif ext in ['.docx', '.doc']:
        return extract_text_from_docx(filepath)
    elif ext in ['.txt', '.rtf']:
        return extract_text_from_txt(filepath)
    return ""

def extract_email(text):
    email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
    matches = re.findall(email_pattern, text)
    if matches:
        return matches[0].strip()
    return ""

def extract_phone(text):
    phone_patterns = [
        r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3,5}\)?[-.\s]?\d{3,5}[-.\s]?\d{3,5}',
        r'\b[6-9]\d{9}\b'  # Standard 10-digit mobile
    ]
    for pattern in phone_patterns:
        matches = re.findall(pattern, text)
        for m in matches:
            cleaned = re.sub(r'[^\d+]', '', m)
            if len(cleaned) >= 10:
                return m.strip()
    return ""

def extract_name(text, filename=""):
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    invalid_keywords = ['curriculum', 'vitae', 'resume', 'profile', 'contact', 'email', 'phone', 'page', 'http', 'github', 'linkedin', 'education', 'skills', 'experience', 'summary', 'objective']
    
    for line in lines[:8]:
        line_clean = re.sub(r'[^a-zA-Z\s]', '', line).strip()
        words = line_clean.split()
        if 2 <= len(words) <= 4:
            if not any(kw in line.lower() for kw in invalid_keywords):
                # Check capitalization
                if all(w[0].isupper() for w in words if len(w) > 1):
                    return line_clean

    # Fallback to filename
    if filename:
        base = os.path.splitext(os.path.basename(filename))[0]
        base = re.sub(r'[-_]', ' ', base)
        base = re.sub(r'(?i)\bresume\b|\bcv\b|\bprofile\b|\bdoc\b', '', base).strip()
        words = [w.capitalize() for w in base.split() if w.isalpha()]
        if words:
            return " ".join(words)

    return "Candidate"

def extract_skills(text, target_skills_list=None):
    found_skills = set()
    all_skills = set(COMMON_SKILLS)
    if target_skills_list:
        for s in target_skills_list:
            if s.strip():
                all_skills.add(s.strip())

    text_lower = text.lower()

    for skill in all_skills:
        s_clean = skill.strip()
        if not s_clean:
            continue
        
        # Regex word boundary check
        # For short or symbol-heavy skills (like C, C++, C#, R, .NET)
        if s_clean in ["C", "R"]:
            pattern = rf'(?<![a-zA-Z0-9]){re.escape(s_clean)}(?![a-zA-Z0-9+#])'
            if re.search(pattern, text):
                found_skills.add(s_clean)
        elif s_clean in ["C++", "C#", ".NET"]:
            pattern = rf'(?<![a-zA-Z0-9]){re.escape(s_clean)}(?![a-zA-Z0-9])'
            if re.search(pattern, text, re.IGNORECASE):
                found_skills.add(s_clean)
        else:
            pattern = rf'\b{re.escape(s_clean.lower())}\b'
            if re.search(pattern, text_lower):
                found_skills.add(s_clean)

    return sorted(list(found_skills), key=lambda s: s.lower())

def extract_education(text):
    edu_keywords = [
        r'\b(?:B\.?Tech|B\.?E|M\.?Tech|M\.?E|B\.?S|M\.?S|B\.?Sc|M\.?Sc|BCA|MCA|Ph\.?D|Bachelor|Master|Diploma|High School)\b',
        r'\b(?:Computer Science|Information Technology|Data Science|Software Engineering|Electrical|Mechanical|Civil)\b',
        r'\b(?:University|Institute|College|Academy|School)\b'
    ]
    
    found_lines = []
    lines = text.split('\n')
    for line in lines:
        line_clean = line.strip()
        if not line_clean:
            continue
        match_count = sum(1 for kw in edu_keywords if re.search(kw, line_clean, re.IGNORECASE))
        if match_count >= 1 and len(line_clean) < 150:
            found_lines.append(line_clean)

    if found_lines:
        return " | ".join(found_lines[:3])
    
    # Generic fallback
    if re.search(r'\b(?:B\.?Tech|Bachelor|B\.?E)\b', text, re.IGNORECASE):
        return "Bachelor's Degree in Engineering / Computer Science"
    elif re.search(r'\b(?:M\.?Tech|Master|M\.?S)\b', text, re.IGNORECASE):
        return "Master's Degree in Computer Science / Technology"
    return "Not specified"

def extract_experience_years(text):
    # Search for explicit patterns like "3 years of experience", "2.5+ yrs exp"
    patterns = [
        r'(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)(?:\s+of)?\s*(?:relevant\s+)?(?:experience|exp|industry\s+experience)',
        r'(?:experience|exp):\s*(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)'
    ]
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE)
        if match:
            try:
                years = float(match.group(1))
                if 0 <= years <= 40:
                    return years
            except ValueError:
                pass

    # Look for year ranges like 2020 - 2023 or 2019 to Present
    year_range_pattern = r'\b(20\d\d|19\d\d)\s*(?:-|–|to)\s*(20\d\d|present|current)\b'
    matches = re.findall(year_range_pattern, text, re.IGNORECASE)
    total_est_years = 0
    from datetime import datetime
    current_year = datetime.now().year

    for start, end in matches:
        try:
            s_yr = int(start)
            e_yr = current_year if end.lower() in ['present', 'current'] else int(end)
            diff = e_yr - s_yr
            if 0 < diff <= 10:
                total_est_years += diff
        except ValueError:
            continue

    if total_est_years > 0:
        return min(float(total_est_years), 25.0)

    # Keyword hints
    if re.search(r'\bsenior\b|\blead\b|\bprincipal\b', text, re.IGNORECASE):
        return 5.0
    elif re.search(r'\bjunior\b|\bintern\b|\bfresher\b|\bgraduate\b', text, re.IGNORECASE):
        return 1.0

    return 1.5

def extract_certifications(text):
    cert_keywords = [
        "AWS Certified", "Solutions Architect", "Developer Associate", "Cloud Practitioner",
        "Google Cloud Certified", "Azure Certified", "Azure Fundamentals", "PMP",
        "Certified Scrum Master", "CSM", "CCNA", "CompTIA", "HackerRank", "Coursera",
        "DeepLearning.AI", "Meta Front-End", "Meta Back-End", "TensorFlow Developer",
        "Oracle Certified", "Python Institute"
    ]
    found = []
    text_lower = text.lower()
    for cert in cert_keywords:
        if cert.lower() in text_lower:
            found.append(cert)
    
    # If explicit Certifications section found
    lines = text.split('\n')
    in_cert_section = False
    for line in lines:
        l = line.strip()
        if re.search(r'^(?:certifications?|licenses?|credentials?)\b', l, re.IGNORECASE):
            in_cert_section = True
            continue
        if in_cert_section:
            if re.search(r'^(?:education|skills|experience|projects|languages)\b', l, re.IGNORECASE):
                in_cert_section = False
            elif l and len(l) < 100:
                clean_l = re.sub(r'^[•\-\*]\s*', '', l)
                if clean_l and clean_l not in found:
                    found.append(clean_l)
                    if len(found) >= 5:
                        break

    return found[:5]

def extract_projects(text):
    projects = []
    lines = text.split('\n')
    in_project_section = False
    for line in lines:
        l = line.strip()
        if re.search(r'^(?:key\s+)?projects?\b', l, re.IGNORECASE):
            in_project_section = True
            continue
        if in_project_section:
            if re.search(r'^(?:education|skills|experience|certifications|contact|summary)\b', l, re.IGNORECASE):
                in_project_section = False
            elif l and len(l) < 120 and (l.startswith('-') or l.startswith('•') or l.startswith('*') or ':' in l):
                clean_p = re.sub(r'^[•\-\*]\s*', '', l)
                if clean_p:
                    projects.append(clean_p)
                    if len(projects) >= 4:
                        break
    return projects

def parse_resume(filepath, original_filename=None):
    filename = original_filename or os.path.basename(filepath)
    raw_text = extract_text(filepath)
    
    name = extract_name(raw_text, filename)
    email = extract_email(raw_text)
    phone = extract_phone(raw_text)
    skills = extract_skills(raw_text)
    education = extract_education(raw_text)
    exp_years = extract_experience_years(raw_text)
    certs = extract_certifications(raw_text)
    projects = extract_projects(raw_text)

    # Basic summary of experience
    exp_summary = f"{exp_years} years estimated experience in software development and technology."
    
    return {
        "name": name,
        "email": email,
        "phone": phone,
        "extracted_skills": skills,
        "education": education,
        "experience_years": exp_years,
        "experience_details": exp_summary,
        "certifications": certs,
        "projects": projects,
        "raw_resume_text": raw_text,
        "resume_filename": filename,
        "resume_filepath": filepath
    }
